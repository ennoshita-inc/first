import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "振り返りシート"

# --- Page Setup: A3 Landscape ---
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = 0.5
ws.page_margins.right = 0.5
ws.page_margins.top = 0.4
ws.page_margins.bottom = 0.4
ws.page_margins.header = 0.2
ws.page_margins.footer = 0.2

# --- Styles ---
title_font = Font(name="游ゴシック", size=18, bold=True, color="FFFFFF")
title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
title_align = Alignment(horizontal="center", vertical="center")

header_font = Font(name="游ゴシック", size=11, bold=True)
header_info_font = Font(name="游ゴシック", size=11)
header_info_align = Alignment(horizontal="left", vertical="center")

section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
section_font = Font(name="游ゴシック", size=11, bold=True)
section_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

boss_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
boss_font = Font(name="游ゴシック", size=11, bold=True)

normal_font = Font(name="游ゴシック", size=10)
normal_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_side = Side(style="thin", color="000000")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

block_header_font = Font(name="游ゴシック", size=12, bold=True, color="FFFFFF")
block_header_fill_left = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
block_header_fill_right = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

footer_font = Font(name="游ゴシック", size=9, color="333333")
footer_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

company_font = Font(name="游ゴシック", size=9, color="666666")
company_align = Alignment(horizontal="right", vertical="center")

sub_label_font = Font(name="游ゴシック", size=9, color="555555")
sub_label_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

# --- Column widths ---
# Columns: A=left section label, B=left sub-label, C-D=left input,
#           E(spacer), F=right section label, G=right sub-label, H-I=right input
col_widths = {
    "A": 5,    # section number
    "B": 22,   # description
    "C": 25,   # input area
    "D": 25,   # input area cont.
    "E": 2,    # spacer
    "F": 5,    # section number
    "G": 22,   # description
    "H": 25,   # input area
    "I": 25,   # input area cont.
}

for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# =====================
# ROW 1: Title
# =====================
row = 1
ws.merge_cells("A1:I1")
ws.row_dimensions[1].height = 42
cell = ws["A1"]
cell.value = "次世代リーダー研修　振り返りシート"
cell.font = title_font
cell.fill = title_fill
cell.alignment = title_align
cell.border = thin_border
# Apply fill/border to merged range
for c in range(2, 10):
    cl = ws.cell(row=1, column=c)
    cl.fill = title_fill
    cl.border = thin_border

# =====================
# ROW 2: Info fields
# =====================
row = 2
ws.row_dimensions[2].height = 28

info_items = [
    ("A2:B2", "氏名："),
    ("C2:D2", "所属："),
    ("F2:G2", "第　　回（　　月　　日実施）"),
    ("H2:I2", "チーム名："),
]
for rng, val in info_items:
    ws.merge_cells(rng)
    cell = ws[rng.split(":")[0]]
    cell.value = val
    cell.font = header_info_font
    cell.alignment = header_info_align
    cell.border = thin_border

# Fill borders for merged cells in row 2
for c in range(1, 10):
    ws.cell(row=2, column=c).border = thin_border

# Spacer column E in row 2
ws.cell(row=2, column=5).border = Border()

# =====================
# ROW 3: Block Headers
# =====================
row = 3
ws.row_dimensions[3].height = 28

ws.merge_cells("A3:D3")
cell = ws["A3"]
cell.value = "研修当日（研修終了時に記入）"
cell.font = block_header_font
cell.fill = block_header_fill_left
cell.alignment = center_align
for c in range(1, 5):
    ws.cell(row=3, column=c).fill = block_header_fill_left
    ws.cell(row=3, column=c).border = thin_border

ws.merge_cells("F3:I3")
cell = ws["F3"]
cell.value = "実践振り返り（次回研修までに記入）"
cell.font = block_header_font
cell.fill = block_header_fill_right
cell.alignment = center_align
for c in range(6, 10):
    ws.cell(row=3, column=c).fill = block_header_fill_right
    ws.cell(row=3, column=c).border = thin_border

# Spacer
ws.cell(row=3, column=5).border = Border()

# =====================
# Helper to write a section block
# =====================
def write_section(start_row, num_col, desc_col, input_start_col, input_end_col,
                  number_text, desc_text, sub_text, height_rows,
                  is_boss=False):
    """Write a section spanning height_rows rows."""
    # Merge section number cell
    if height_rows > 1:
        ws.merge_cells(
            start_row=start_row, start_column=num_col,
            end_row=start_row + height_rows - 1, end_column=num_col
        )
    cell_num = ws.cell(row=start_row, column=num_col)
    cell_num.value = number_text
    fill = boss_fill if is_boss else section_fill
    cell_num.fill = fill
    cell_num.font = section_font if not is_boss else boss_font
    cell_num.alignment = section_align
    cell_num.border = thin_border

    # Merge description cell
    if height_rows > 1:
        ws.merge_cells(
            start_row=start_row, start_column=desc_col,
            end_row=start_row + height_rows - 1, end_column=desc_col
        )
    cell_desc = ws.cell(row=start_row, column=desc_col)
    cell_desc.value = sub_text
    cell_desc.fill = fill
    cell_desc.font = sub_label_font if not is_boss else Font(name="游ゴシック", size=9, color="555555", bold=True)
    cell_desc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell_desc.border = thin_border

    # Merge input area
    if height_rows > 1 or input_end_col > input_start_col:
        ws.merge_cells(
            start_row=start_row, start_column=input_start_col,
            end_row=start_row + height_rows - 1, end_column=input_end_col
        )
    cell_input = ws.cell(row=start_row, column=input_start_col)
    cell_input.value = ""
    cell_input.font = normal_font
    cell_input.alignment = normal_align
    cell_input.border = thin_border

    # Apply borders and fills to all cells in the range
    for r in range(start_row, start_row + height_rows):
        ws.cell(row=r, column=num_col).fill = fill
        ws.cell(row=r, column=desc_col).fill = fill
        for c in [num_col, desc_col, input_start_col, input_end_col]:
            ws.cell(row=r, column=c).border = thin_border
        # Fill input cells borders
        for c in range(input_start_col, input_end_col + 1):
            ws.cell(row=r, column=c).border = thin_border

    # Set row heights
    row_h = 20
    if height_rows >= 5:
        row_h = 28
    elif height_rows >= 3:
        row_h = 25
    for r in range(start_row, start_row + height_rows):
        ws.row_dimensions[r].height = row_h


def write_single_line(row_num, num_col, desc_col, input_start_col, input_end_col,
                      label_text, value_text, is_boss=False):
    """Write a single-line row (e.g., date, signature)."""
    fill = boss_fill
    ws.cell(row=row_num, column=num_col).value = ""
    ws.cell(row=row_num, column=num_col).fill = fill
    ws.cell(row=row_num, column=num_col).border = thin_border

    ws.cell(row=row_num, column=desc_col).value = label_text
    ws.cell(row=row_num, column=desc_col).fill = fill
    ws.cell(row=row_num, column=desc_col).font = Font(name="游ゴシック", size=10, bold=True)
    ws.cell(row=row_num, column=desc_col).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_num, column=desc_col).border = thin_border

    ws.merge_cells(
        start_row=row_num, start_column=input_start_col,
        end_row=row_num, end_column=input_end_col
    )
    ws.cell(row=row_num, column=input_start_col).value = value_text
    ws.cell(row=row_num, column=input_start_col).font = normal_font
    ws.cell(row=row_num, column=input_start_col).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_num, column=input_start_col).border = thin_border
    ws.cell(row=row_num, column=input_end_col).border = thin_border

    ws.row_dimensions[row_num].height = 22


# =====================
# LEFT BLOCK sections (columns A-D, starting row 4)
# =====================
cur_row = 4

# ① 本日の研修テーマ (2 rows)
write_section(cur_row, 1, 2, 3, 4, "①", "本日の研修テーマ", "（事務局記入）", 2)
cur_row += 2

# ② 学んだこと (5 rows)
write_section(cur_row, 1, 2, 3, 4, "②", "学んだこと", "研修で学んだ内容を\n具体的に記録", 5)
cur_row += 5

# ③ 気づき (5 rows)
write_section(cur_row, 1, 2, 3, 4, "③", "気づき", "学びから得た\n気づきを記録", 5)
cur_row += 5

# ④ 活用計画 (5 rows)
write_section(cur_row, 1, 2, 3, 4, "④", "活用計画", "日常・現場の仕事に\nどう活かすか", 5)
cur_row += 5

# ⑤ 上司コメント (3 rows)
write_section(cur_row, 1, 2, 3, 4, "⑤", "上司コメント", "上司（コーディネーター）が\n計画を確認しコメント", 3, is_boss=True)
cur_row += 3

# 上司確認日
write_single_line(cur_row, 1, 2, 3, 4, "上司確認日", "　　　年　　月　　日", is_boss=True)
cur_row += 1

# 上司署名
write_single_line(cur_row, 1, 2, 3, 4, "上司署名", "", is_boss=True)
left_end_row = cur_row

# =====================
# RIGHT BLOCK sections (columns F-I, starting row 4)
# =====================
cur_row = 4

# ⑥ 実践したこと (5 rows)
write_section(cur_row, 6, 7, 8, 9, "⑥", "実践したこと", "活用計画に基づき\n実際に行ったことを記録", 5)
cur_row += 5

# ⑦ 実践からの気づき・教訓 (5 rows)
write_section(cur_row, 6, 7, 8, 9, "⑦", "実践からの\n気づき・教訓", "実践から得た\n気づきや教訓を記録", 5)
cur_row += 5

# ⑧ 次回に向けて (3 rows)
write_section(cur_row, 6, 7, 8, 9, "⑧", "次回に向けて", "次に取り組みたいこと\n・改善点", 3)
cur_row += 3

# ⑨ 上司コメント (3 rows)
write_section(cur_row, 6, 7, 8, 9, "⑨", "上司コメント", "上司（コーディネーター）が\n振り返りを確認しコメント", 3, is_boss=True)
cur_row += 3

# 上司確認日
write_single_line(cur_row, 6, 7, 8, 9, "上司確認日", "　　　年　　月　　日", is_boss=True)
cur_row += 1

# 上司署名
write_single_line(cur_row, 6, 7, 8, 9, "上司署名", "", is_boss=True)
right_end_row = cur_row

# =====================
# Spacer column E - clear borders for all data rows
# =====================
for r in range(3, max(left_end_row, right_end_row) + 1):
    ws.cell(row=r, column=5).border = Border()

# =====================
# FOOTER
# =====================
footer_row = max(left_end_row, right_end_row) + 1
ws.row_dimensions[footer_row].height = 32
ws.merge_cells(f"A{footer_row}:I{footer_row}")
cell = ws.cell(row=footer_row, column=1)
cell.value = "※研修終了後、左ブロックを記入 → 上司に提出・確認 → 次回研修までに右ブロックを記入 → 上司に提出・確認 → 次回研修に持参"
cell.font = footer_font
cell.alignment = footer_align
for c in range(1, 10):
    ws.cell(row=footer_row, column=c).border = Border()

company_row = footer_row + 1
ws.row_dimensions[company_row].height = 20
ws.merge_cells(f"A{company_row}:I{company_row}")
cell = ws.cell(row=company_row, column=1)
cell.value = "株式会社えんのした"
cell.font = company_font
cell.alignment = company_align

# =====================
# Print area
# =====================
ws.print_area = f"A1:I{company_row}"

# Save
wb.save("/home/user/first/review-sheet.xlsx")
print("Done!")
