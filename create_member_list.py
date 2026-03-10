import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# --------------- Common Styles ---------------
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="游ゴシック", bold=True, color="FFFFFF", size=11)
title_font = Font(name="游ゴシック", bold=True, size=16)
subtitle_font = Font(name="游ゴシック", size=11, color="555555")
note_font = Font(name="游ゴシック", size=10)
note_title_font = Font(name="游ゴシック", size=10, bold=True)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
note_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Tab colors
TAB_GREEN = "00B050"
TAB_BLUE = "4472C4"
TAB_ORANGE = "ED7D31"
TAB_PURPLE = "7030A0"


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_align


def style_data_area(ws, start_row, end_row, cols, center_cols=None):
    if center_cols is None:
        center_cols = {1}
    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = 30
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.alignment = center_align if c in center_cols else left_align


def add_note(ws, start_row, note_text, merge_width=6):
    ws.merge_cells(
        start_row=start_row, start_column=1,
        end_row=start_row + 5, end_column=merge_width,
    )
    cell = ws.cell(row=start_row, column=1)
    cell.value = note_text
    cell.font = note_font
    cell.alignment = note_align


NOTE_TEXT = (
    "【研修の役割構成】\n"
    "・塾生：部長候補（次世代リーダー）最大12名 → 全回参加\n"
    "・コーディネーター：現職部長クラス 4名 → 各チームの支援・伴走\n"
    "・塾長/オブザーバー：経営層 1〜2名 → 可能な回にオブザーブ参加\n"
    "※塾生が15名を超えるとフォローや発表に時間がかかるため、最大12名を推奨いたします。"
)

# =========================================================
# Sheet 1: 塾生名簿
# =========================================================
ws1 = wb.active
ws1.title = "塾生名簿"
ws1.sheet_properties.tabColor = TAB_GREEN

headers1 = ["No.", "氏名", "年齢", "役職", "所属部署", "勤続年数", "担当業務", "スキル・特徴", "PCスキル（PowerPoint等）", "備考"]
widths1 = [6, 18, 8, 15, 18, 10, 22, 35, 25, 20]
num_cols1 = len(headers1)

# Title
ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols1)
c = ws1.cell(row=1, column=1, value="次世代リーダー研修 塾生名簿")
c.font = title_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[1].height = 35

# Subtitle
ws1.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols1)
c = ws1.cell(row=2, column=1, value="※最大12名。下記の項目をご記入ください。")
c.font = subtitle_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[2].height = 25

# Header row
header_row1 = 4
for i, h in enumerate(headers1, 1):
    ws1.cell(row=header_row1, column=i, value=h)
style_header(ws1, header_row1, num_cols1)
ws1.row_dimensions[header_row1].height = 35

# Data rows (12)
data_start1 = header_row1 + 1
data_end1 = data_start1 + 11
for r in range(data_start1, data_end1 + 1):
    ws1.cell(row=r, column=1, value=r - data_start1 + 1)
style_data_area(ws1, data_start1, data_end1, num_cols1, center_cols={1, 3, 6})

# Column widths
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

# Note
add_note(ws1, data_end1 + 2, NOTE_TEXT, merge_width=num_cols1)

# =========================================================
# Sheet 2: コーディネーター名簿
# =========================================================
ws2 = wb.create_sheet("コーディネーター名簿")
ws2.sheet_properties.tabColor = TAB_BLUE

headers2 = ["No.", "氏名", "年齢", "役職", "所属部署", "担当チーム", "コーディネーターとしての役割・期待", "備考"]
widths2 = [6, 18, 8, 15, 18, 15, 40, 20]
num_cols2 = len(headers2)

ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols2)
c = ws2.cell(row=1, column=1, value="次世代リーダー研修 コーディネーター名簿")
c.font = title_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[1].height = 35

ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols2)
c = ws2.cell(row=2, column=1, value="※各チーム1名（4チーム×1名＝4名）。塾生の上司クラスの方。")
c.font = subtitle_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[2].height = 25

header_row2 = 4
for i, h in enumerate(headers2, 1):
    ws2.cell(row=header_row2, column=i, value=h)
style_header(ws2, header_row2, num_cols2)
ws2.row_dimensions[header_row2].height = 35

data_start2 = header_row2 + 1
data_end2 = data_start2 + 3
for r in range(data_start2, data_end2 + 1):
    ws2.cell(row=r, column=1, value=r - data_start2 + 1)
style_data_area(ws2, data_start2, data_end2, num_cols2, center_cols={1, 3})

for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

add_note(ws2, data_end2 + 2, NOTE_TEXT, merge_width=num_cols2)

# =========================================================
# Sheet 3: 塾長・オブザーバー
# =========================================================
ws3 = wb.create_sheet("塾長・オブザーバー")
ws3.sheet_properties.tabColor = TAB_ORANGE

headers3 = ["No.", "氏名", "役職", "参加形態（塾長/オブザーバー）", "参加予定頻度", "備考"]
widths3 = [6, 18, 18, 30, 20, 25]
num_cols3 = len(headers3)

ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols3)
c = ws3.cell(row=1, column=1, value="次世代リーダー研修 塾長・オブザーバー")
c.font = title_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws3.row_dimensions[1].height = 35

ws3.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols3)
c = ws3.cell(row=2, column=1, value="※経営層・役員クラスの方。可能な回にオブザーブ参加。")
c.font = subtitle_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws3.row_dimensions[2].height = 25

header_row3 = 4
for i, h in enumerate(headers3, 1):
    ws3.cell(row=header_row3, column=i, value=h)
style_header(ws3, header_row3, num_cols3)
ws3.row_dimensions[header_row3].height = 35

data_start3 = header_row3 + 1
data_end3 = data_start3 + 1
for r in range(data_start3, data_end3 + 1):
    ws3.cell(row=r, column=1, value=r - data_start3 + 1)
style_data_area(ws3, data_start3, data_end3, num_cols3, center_cols={1})

for i, w in enumerate(widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

add_note(ws3, data_end3 + 2, NOTE_TEXT, merge_width=num_cols3)

# =========================================================
# Sheet 4: チーム編成表
# =========================================================
ws4 = wb.create_sheet("チーム編成表")
ws4.sheet_properties.tabColor = TAB_PURPLE

headers4 = ["チーム名", "コーディネーター", "メンバー1", "メンバー2", "メンバー3"]
widths4 = [18, 20, 20, 20, 20]
num_cols4 = len(headers4)

ws4.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols4)
c = ws4.cell(row=1, column=1, value="チーム編成表")
c.font = title_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws4.row_dimensions[1].height = 35

ws4.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols4)
c = ws4.cell(row=2, column=1, value="※4チーム×3名＋コーディネーター1名の構成")
c.font = subtitle_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws4.row_dimensions[2].height = 25

header_row4 = 4
for i, h in enumerate(headers4, 1):
    ws4.cell(row=header_row4, column=i, value=h)
style_header(ws4, header_row4, num_cols4)
ws4.row_dimensions[header_row4].height = 35

data_start4 = header_row4 + 1
data_end4 = data_start4 + 3
team_names = ["チームA", "チームB", "チームC", "チームD"]
for idx, r in enumerate(range(data_start4, data_end4 + 1)):
    ws4.cell(row=r, column=1, value=team_names[idx])
style_data_area(ws4, data_start4, data_end4, num_cols4, center_cols={1})

for i, w in enumerate(widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

add_note(ws4, data_end4 + 2, NOTE_TEXT, merge_width=num_cols4)

# =========================================================
# Save
# =========================================================
wb.save("/home/user/first/member-list.xlsx")
print("Done: /home/user/first/member-list.xlsx")
